#!/usr/bin/env python3
"""读取docx内容，定位'考核等级'单元格"""
import os

import glob
files = glob.glob("/Users/mac/Desktop/OA-EPP/成信教务函*名师*申报表*.docx")
if not files:
    print("ERROR: file not found")
    exit(1)
docx_path = files[0]
print(f"Found: {docx_path}")

from docx import Document
doc = Document(docx_path)

print("=== 段落 ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        print(f"[P{i}][{p.style.name}] {t}")

print("\n=== 表格 ===")
for ti, table in enumerate(doc.tables):
    print(f"\n--- 表格 {ti} ({len(table.rows)}行 x {len(table.columns)}列) ---")
    for ri, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        print(f"  行{ri}: {cells}")
